import base64
import tempfile
from pathlib import Path
from unittest.mock import patch
from io import BytesIO

import docx

from app.models.system import GmailScannedMessage, ResumeIngestion
from app.services import gmail_inbox


def make_docx():
    document = docx.Document()
    document.add_paragraph("Alex Example")
    document.add_paragraph("alex@example.test | 555-010-1234")
    document.add_paragraph("Professional Experience")
    document.add_paragraph("Backend Engineer | Example Systems")
    document.add_paragraph("Built production APIs and PostgreSQL services.")
    document.add_paragraph("2021 - Present")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class FakeDb:
    def __init__(self):
        self.rows = []
        self.scanned = []
        self.pending = []

    def scalar(self, query):
        if "gmail_scanned_messages" in str(query):
            return 1 if self.scanned else None
        return 1 if self.rows else None

    def add(self, row):
        self.pending.append(row)

    def commit(self):
        for item in self.pending:
            if isinstance(item, GmailScannedMessage):
                self.scanned.append(item)
            elif isinstance(item, ResumeIngestion):
                self.rows.append(item)
        self.pending.clear()

    def rollback(self):
        self.pending.clear()


attachment = make_docx()
payload = {
    "filename": "Alex-Example.docx",
    "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "body": {"data": base64.urlsafe_b64encode(attachment).decode().rstrip("=")},
}
message = {
    "snippet": "A new applicant submitted their application",
    "payload": {
        "headers": [{"name": "Subject", "value": "New applicant for Backend Engineer"}],
        "parts": [payload],
    },
}

with tempfile.TemporaryDirectory() as directory:
    db = FakeDb()
    with (
        patch.object(gmail_inbox.settings, "gmail_client_id", "client"),
        patch.object(gmail_inbox.settings, "gmail_client_secret", "secret"),
        patch.object(gmail_inbox.settings, "gmail_refresh_token", "refresh"),
        patch.object(gmail_inbox.settings, "storage_dir", Path(directory)),
        patch.object(gmail_inbox, "_access_token", return_value="access"),
        patch.object(gmail_inbox, "_gmail_json", side_effect=[
            {"messages": [{"id": "message-123"}]}, message,
            {"messages": [{"id": "message-123"}]},
        ]),
    ):
        assert gmail_inbox.poll_gmail(db) == 1
        assert gmail_inbox.poll_gmail(db) == 0

    assert len(db.rows) == 1
    row = db.rows[0]
    assert row.source == "email"
    assert row.status == "Received"
    assert row.original_filename == "Alex-Example.docx"
    assert (Path(directory) / row.staged_path).read_bytes() == attachment

assert gmail_inbox._extension(
    {"filename": "document.pdf", "mimeType": "application/pdf"}, b"%PDF-1.7 data"
) == "pdf"
assert gmail_inbox._extension(
    {"filename": "not-a-resume.pdf", "mimeType": "application/pdf"}, b"not a PDF"
) is None
assert gmail_inbox._is_resume("Email: a@example.test\nSkills\nWork history", "file.pdf", "")
assert not gmail_inbox._is_resume("A generic document", "document.pdf", "Invoice attached")
assert not gmail_inbox._is_resume("Invoice payment due", "invoice.pdf", "New application submitted")
assert gmail_inbox._is_resume("", "candidate-resume.pdf", "New application submitted")
print("Gmail resume inbox checks passed.")
