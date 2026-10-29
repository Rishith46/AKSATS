import tempfile
from pathlib import Path

import docx

from app.services.text_extract import ExtractionError, extract_text

with tempfile.TemporaryDirectory() as d:
    p = Path(d) / "mock_resume.docx"
    doc = docx.Document()
    doc.add_paragraph("MOCK PERSON (test data, not a real candidate)")
    doc.add_paragraph("mock.person@example.com | 555-010-0000 | Dallas, TX")
    doc.add_paragraph("Senior Java Developer, Acme Mock Corp, 2021 - Present")
    doc.add_paragraph("Built Java microservices with Spring Boot on AWS.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Education"
    table.cell(0, 1).text = "Example University"
    doc.sections[0].header.paragraphs[0].text = "Header detail"
    doc.sections[0].footer.paragraphs[0].text = "Footer detail"
    doc.save(str(p))
    text = extract_text(p)
    print("DOCX extracted characters:", len(text))
    assert "mock.person@example.com" in text
    assert "Spring Boot" in text
    assert "Example University" in text
    assert "Header detail" in text
    assert "Footer detail" in text

    bad = Path(d) / "empty.docx"
    docx.Document().save(str(bad))
    try:
        extract_text(bad)
        print("Empty file: NOT rejected (unexpected)")
    except ExtractionError as e:
        print("Empty file rejected:", e)

    short = Path(d) / "short.docx"
    short_doc = docx.Document()
    short_doc.add_paragraph("Ada Lovelace")
    short_doc.save(str(short))
    assert extract_text(short) == "Ada Lovelace"
    print("Short DOCX accepted")