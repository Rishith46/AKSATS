import hashlib
import uuid

import docx
from sqlalchemy import delete, select

from app.config import settings
from app.db import SessionLocal
from app.models.candidates import (Candidate, CandidateDocument, CandidateExperience,
                                   CandidateSkill, Skill)
from app.services.parse_pipeline import parse_document

MOCK = [
    "Jane Mockperson",
    "Dallas, TX | jane.mock@example.com | (555) 010-0123",
    "EXPERIENCE",
    "Senior Java Developer - Acme Mock Corp - Jan 2021 to Present",
    "Built Java microservices with Spring Boot on AWS.",
    "Java Developer - Beta Test Systems - 2018 - 2020",
    "Maintained internal Java applications.",
    "EDUCATION",
    "BS Computer Science, Example University, 2017",
]

with SessionLocal() as db:
    c = Candidate(full_name=None)
    db.add(c)
    db.flush()
    rel = f"candidates/{c.id}/{uuid.uuid4().hex}.docx"
    full = settings.storage_dir / rel
    full.parent.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    for line in MOCK:
        d.add_paragraph(line)
    d.save(str(full))
    data = full.read_bytes()
    doc = CandidateDocument(candidate_id=c.id, kind="resume", original_filename="MOCK.docx",
                            storage_path=rel, size_bytes=len(data),
                            sha256=hashlib.sha256(data).hexdigest(),
                            version_no=1, is_current=True)
    db.add(doc)
    db.commit()
    cid = c.id

    try:
        print("Result:", parse_document(db, doc.id))
        db.expire_all()
        c = db.get(Candidate, cid)
        print("Candidate:", c.full_name, "|", c.email, "|", c.phone, "|", c.location,
              "|", c.parse_status)
        print("field_sources:", c.field_sources)
        for e in db.scalars(select(CandidateExperience).where(CandidateExperience.candidate_id == cid)):
            print("Experience:", e.company, "|", e.title, "|", e.start_date, e.start_precision, "|", e.end_date, e.end_precision, "|", e.is_current)
        names = [n for (n,) in db.execute(select(Skill.name).join(CandidateSkill, CandidateSkill.skill_id == Skill.id)
                                          .where(CandidateSkill.candidate_id == cid)).all()]
        print("Skills:", sorted(names))
    finally:
        db.delete(db.get(Candidate, cid))
        db.commit()
        full.unlink(missing_ok=True)
        try:
            full.parent.rmdir()
        except OSError:
            pass
        print("Mock candidate and file removed.")