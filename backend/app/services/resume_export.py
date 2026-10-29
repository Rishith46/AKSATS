import re
from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether, Paragraph, SimpleDocTemplate, Spacer,
)

from app.services.experience_edit import format_partial_date


def _text(value):
    return str(value).strip() if value is not None and str(value).strip() else None


def _date_range(experience):
    start = format_partial_date(experience.start_date, experience.start_precision)
    end = "Present" if experience.is_current else format_partial_date(
        experience.end_date, experience.end_precision)
    return " to ".join(value for value in (start, end) if value) or None


def _skill_label(skill):
    details = []
    if skill.years is not None:
        details.append(f"{float(skill.years):g} years")
    if skill.last_used_year is not None:
        details.append(f"last used {skill.last_used_year}")
    return f"{skill.name} ({', '.join(details)})" if details else skill.name


def _add_docx_item(document, parts, description=None):
    title = " | ".join(part for part in parts if part)
    if title:
        document.add_paragraph(title, style="List Bullet")
    if description:
        document.add_paragraph(description)


def build_docx(candidate, experiences, education, certifications, skills):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    document.add_heading(_text(candidate.full_name) or "Candidate Resume", 0)
    contact = [
        item for item in (
            _text(candidate.email), _text(candidate.phone), _text(candidate.location),
            _text(candidate.linkedin_url),
        ) if item
    ]
    if contact:
        document.add_paragraph(" | ".join(contact))

    document.add_heading("Professional Experience", level=1)
    for exp in experiences:
        _add_docx_item(document, [
            _text(exp.title), _text(exp.company), _text(exp.location), _date_range(exp),
        ], _text(exp.description))

    document.add_heading("Skills", level=1)
    for skill in skills:
        _add_docx_item(document, [_skill_label(skill)], _text(skill.evidence))

    if education:
        document.add_heading("Education", level=1)
        for item in education:
            _add_docx_item(document, [
                _text(item.degree), _text(item.field_of_study), _text(item.institution),
                " | ".join(str(year) for year in (item.start_year, item.end_year) if year) or None,
            ])

    if certifications:
        document.add_heading("Certifications", level=1)
        for item in certifications:
            _add_docx_item(document, [
                _text(item.name), _text(item.issuer), str(item.year) if item.year else None,
            ])

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _paragraph(value, style):
    return Paragraph(escape(value).replace("\n", "<br/>"), style)


def build_pdf(candidate, experiences, education, certifications, skills):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ResumeName", parent=styles["Title"], alignment=TA_CENTER,
        fontName="Helvetica-Bold", fontSize=20, leading=24, spaceAfter=6,
    ))
    styles.add(ParagraphStyle(
        name="ResumeSection", parent=styles["Heading2"], textColor=colors.HexColor("#243746"),
        spaceBefore=10, spaceAfter=4,
    ))
    styles.add(ParagraphStyle(
        name="ResumeBody", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=9.5, leading=13, spaceAfter=3,
    ))
    styles.add(ParagraphStyle(
        name="ResumeSmall", parent=styles["BodyText"], fontSize=8.5, leading=11,
        textColor=colors.HexColor("#555555"), spaceAfter=3,
    ))
    story = [_paragraph(_text(candidate.full_name) or "Candidate Resume", styles["ResumeName"])]
    contact = [item for item in (
        _text(candidate.email), _text(candidate.phone), _text(candidate.location),
        _text(candidate.linkedin_url),
    ) if item]
    if contact:
        story.append(_paragraph(" | ".join(contact), styles["ResumeSmall"]))

    def heading(value):
        story.append(Paragraph(value, styles["ResumeSection"]))

    heading("Professional Experience")
    for exp in experiences:
        title = " | ".join(part for part in (
            _text(exp.title), _text(exp.company), _text(exp.location), _date_range(exp),
        ) if part)
        group = []
        if title:
            group.append(_paragraph(title, styles["ResumeBody"]))
        if exp.description:
            group.append(_paragraph(exp.description, styles["ResumeSmall"]))
        if group:
            story.append(KeepTogether(group))

    heading("Skills")
    for skill in skills:
        story.append(_paragraph(_skill_label(skill), styles["ResumeBody"]))
        if skill.evidence:
            story.append(_paragraph(skill.evidence, styles["ResumeSmall"]))

    if education:
        heading("Education")
        for item in education:
            parts = [
                _text(item.degree), _text(item.field_of_study), _text(item.institution),
                " | ".join(str(year) for year in (item.start_year, item.end_year) if year) or None,
            ]
            line = " | ".join(part for part in parts if part)
            if line:
                story.append(_paragraph(line, styles["ResumeBody"]))

    if certifications:
        heading("Certifications")
        for item in certifications:
            line = " | ".join(part for part in (
                _text(item.name), _text(item.issuer), str(item.year) if item.year else None,
            ) if part)
            if line:
                story.append(_paragraph(line, styles["ResumeBody"]))

    output = BytesIO()
    SimpleDocTemplate(
        output, pagesize=letter, rightMargin=0.65 * inch, leftMargin=0.65 * inch,
        topMargin=0.55 * inch, bottomMargin=0.55 * inch,
    ).build(story)
    return output.getvalue()


def filename_stem(name):
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", name or "candidate").strip("-._")
    return stem or "candidate"
